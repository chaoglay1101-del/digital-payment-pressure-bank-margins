import pandas as pd
import numpy as np
from linearmodels.panel import PanelOLS
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# =====================================================
# 1. Load merged bank panel with digital payment pressure
# =====================================================
df = pd.read_csv("bank_panel_with_digital_payment_pressure.csv")

# =====================================================
# 2. Create control variables
# =====================================================
df["Size"] = np.log(df["at"])
df["Credit_Risk"] = df["pcl"] / df["at"]
df["Interest_Expense_Ratio"] = df["xint"] / df["at"]
df["Deposit_Ratio"] = df["dptc"] / df["at"]

# Optional: log digital pressure if values are very large
# Use log(1+x) to avoid problems with zeros
df["Log_Digital_Payment_Pressure"] = np.log1p(df["Digital_Payment_Pressure"])

# =====================================================
# 3. Set panel index
# =====================================================
df = df.set_index(["gvkey", "fyear"])

# =====================================================
# 4. Define regression variables
# =====================================================
reg_vars = [
    "Net_Fee_Margin",
    "Digital_Payment_Pressure",
    "Log_Digital_Payment_Pressure",
    "Capital_Adequacy_Proxy",
    "Size",
    "Credit_Risk",
    "Interest_Expense_Ratio",
    "Deposit_Ratio"
]

reg_df = df[reg_vars].dropna()

print("Regression observations:", len(reg_df))
print("Number of banks:", reg_df.index.get_level_values("gvkey").nunique())

# =====================================================
# 5. Model specifications
# =====================================================
models = {
    "Model 1": (
        "Net_Fee_Margin ~ 1 + Digital_Payment_Pressure "
        "+ EntityEffects + TimeEffects"
    ),
    "Model 2": (
        "Net_Fee_Margin ~ 1 + Digital_Payment_Pressure "
        "+ Capital_Adequacy_Proxy + Size + Credit_Risk "
        "+ EntityEffects + TimeEffects"
    ),
    "Model 3": (
        "Net_Fee_Margin ~ 1 + Digital_Payment_Pressure "
        "+ Capital_Adequacy_Proxy + Size + Credit_Risk "
        "+ Interest_Expense_Ratio + Deposit_Ratio "
        "+ EntityEffects + TimeEffects"
    )
}

results = {}

for name, formula in models.items():
    print("\n" + "=" * 80)
    print(name)
    print("=" * 80)

    model = PanelOLS.from_formula(formula, data=reg_df)
    result = model.fit(cov_type="clustered", cluster_entity=True)

    results[name] = result
    print(result)

# =====================================================
# 6. Helper functions for Word table
# =====================================================
def stars(p_value):
    if p_value < 0.01:
        return "***"
    elif p_value < 0.05:
        return "**"
    elif p_value < 0.10:
        return "*"
    else:
        return ""

def coef_and_se(result, variable):
    if variable in result.params.index:
        coef = result.params[variable]
        se = result.std_errors[variable]
        p = result.pvalues[variable]
        return f"{coef:.4f}{stars(p)}", f"({se:.4f})"
    return "", ""

def add_centered_paragraph(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(10)

# =====================================================
# 7. Create Word document
# =====================================================
doc = Document()

doc.add_heading("Main Regression Results", level=1)

doc.add_paragraph(
    "This table reports the fixed effects regression results examining the relationship "
    "between digital payment pressure and banks' net fee margin. Digital payment pressure "
    "is measured using the IMF Financial Access Survey indicator for mobile and internet "
    "banking transactions by commercial banks during the reference year. Standard errors "
    "clustered at the bank level are reported in parentheses."
)

doc.add_heading("Table X. Digital Payment Pressure and Net Fee Margin", level=2)

# Display variables
var_map = {
    "Digital Payment Pressure": "Digital_Payment_Pressure",
    "Capital Adequacy Proxy": "Capital_Adequacy_Proxy",
    "Size": "Size",
    "Credit Risk": "Credit_Risk",
    "Interest Expense Ratio": "Interest_Expense_Ratio",
    "Deposit Ratio": "Deposit_Ratio",
}

rows = [
    "Digital Payment Pressure",
    "",
    "Capital Adequacy Proxy",
    "",
    "Size",
    "",
    "Credit Risk",
    "",
    "Interest Expense Ratio",
    "",
    "Deposit Ratio",
    "",
    "Bank Fixed Effects",
    "Year Fixed Effects",
    "Clustered Standard Errors",
    "Observations",
    "Banks",
    "R-squared",
    "Within R-squared",
    "Robust F-stat p-value",
]

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"

headers = table.rows[0].cells
headers[0].text = "Variables"
headers[1].text = "Model 1"
headers[2].text = "Model 2"
headers[3].text = "Model 3"

last_variable_label = None

for label in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = label

    if label in var_map:
        last_variable_label = label

    for i, model_name in enumerate(["Model 1", "Model 2", "Model 3"], start=1):
        result = results[model_name]

        if label in var_map:
            coef, se = coef_and_se(result, var_map[label])
            add_centered_paragraph(row_cells[i], coef)

        elif label == "":
            if last_variable_label in var_map:
                coef, se = coef_and_se(result, var_map[last_variable_label])
                add_centered_paragraph(row_cells[i], se)
            else:
                add_centered_paragraph(row_cells[i], "")

        elif label == "Bank Fixed Effects":
            add_centered_paragraph(row_cells[i], "Yes")

        elif label == "Year Fixed Effects":
            add_centered_paragraph(row_cells[i], "Yes")

        elif label == "Clustered Standard Errors":
            add_centered_paragraph(row_cells[i], "Yes")

        elif label == "Observations":
            add_centered_paragraph(row_cells[i], str(int(result.nobs)))

        elif label == "Banks":
            n_banks = reg_df.index.get_level_values("gvkey").nunique()
            add_centered_paragraph(row_cells[i], str(n_banks))

        elif label == "R-squared":
            add_centered_paragraph(row_cells[i], f"{result.rsquared:.4f}")

        elif label == "Within R-squared":
            add_centered_paragraph(row_cells[i], f"{result.rsquared_within:.4f}")

        elif label == "Robust F-stat p-value":
            add_centered_paragraph(row_cells[i], f"{result.f_statistic_robust.pval:.4f}")

# Notes
doc.add_paragraph(
    "Notes: Standard errors clustered at the bank level are reported in parentheses. "
    "*** p < 0.01, ** p < 0.05, * p < 0.10. "
    "The dependent variable is Net Fee Margin. "
    "All models include bank and year fixed effects."
)

doc.add_heading("Preliminary Interpretation", level=2)

# Auto-generate simple interpretation based on Model 3
main_coef = results["Model 3"].params.get("Digital_Payment_Pressure", np.nan)
main_p = results["Model 3"].pvalues.get("Digital_Payment_Pressure", np.nan)

if pd.notna(main_coef) and pd.notna(main_p):
    direction = "positive" if main_coef > 0 else "negative"
    significance = (
        "statistically significant" if main_p < 0.10 else "statistically insignificant"
    )

    doc.add_paragraph(
        f"In the full specification, the coefficient of Digital Payment Pressure is "
        f"{direction} and {significance}. This suggests that digital payment pressure "
        f"is {'associated with' if main_p < 0.10 else 'not strongly associated with'} "
        f"banks' net fee margin after controlling for bank-level financial characteristics, "
        f"bank fixed effects, and year fixed effects."
    )

doc.save("main_regression_results.docx")

print("\nWord file exported successfully: main_regression_results.docx")