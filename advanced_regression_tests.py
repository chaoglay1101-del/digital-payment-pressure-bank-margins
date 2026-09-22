import pandas as pd
import numpy as np
from linearmodels.panel import PanelOLS
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# =====================================================
# 1. Load data
# =====================================================
df = pd.read_csv("bank_panel_with_digital_payment_pressure.csv")

# =====================================================
# 2. Create variables
# =====================================================
df["Size"] = np.log(df["at"])
df["Credit_Risk"] = df["pcl"] / df["at"]
df["Interest_Expense_Ratio"] = df["xint"] / df["at"]
df["Deposit_Ratio"] = df["dptc"] / df["at"]

# Digital payment pressure is already measured as % of GDP.
# Use log form for easier interpretation and to avoid coefficients showing as 0.0000.
df["Log_Digital_Payment_Pressure"] = np.log1p(df["Digital_Payment_Pressure"])

# Squared term for non-linear test
df["Log_Digital_Payment_Pressure_Sq"] = df["Log_Digital_Payment_Pressure"] ** 2

# Large bank dummy based on median Size
# Large_Bank = 1 if the bank-year observation has Size above sample median
median_size = df["Size"].median()
df["Large_Bank"] = np.where(df["Size"] >= median_size, 1, 0)

# Interaction term
df["Digital_x_LargeBank"] = df["Log_Digital_Payment_Pressure"] * df["Large_Bank"]

# =====================================================
# 3. Set panel index
# =====================================================
df = df.set_index(["gvkey", "fyear"])

# =====================================================
# 4. Select regression variables
# =====================================================
reg_vars = [
    "Net_Fee_Margin",
    "Log_Digital_Payment_Pressure",
    "Log_Digital_Payment_Pressure_Sq",
    "Large_Bank",
    "Digital_x_LargeBank",
    "Capital_Adequacy_Proxy",
    "Size",
    "Credit_Risk",
    "Interest_Expense_Ratio",
    "Deposit_Ratio"
]

reg_df = df[reg_vars].dropna()

print("Regression observations:", len(reg_df))
print("Number of banks:", reg_df.index.get_level_values("gvkey").nunique())

# =====================================================
# 5. Model specifications
# =====================================================
models = {
    "Model 1 Baseline": (
        "Net_Fee_Margin ~ 1 + Log_Digital_Payment_Pressure "
        "+ Capital_Adequacy_Proxy + Size + Credit_Risk "
        "+ Interest_Expense_Ratio + Deposit_Ratio "
        "+ EntityEffects + TimeEffects"
    ),

    "Model 2 Interaction": (
        "Net_Fee_Margin ~ 1 + Log_Digital_Payment_Pressure "
        "+ Digital_x_LargeBank "
        "+ Capital_Adequacy_Proxy + Size + Credit_Risk "
        "+ Interest_Expense_Ratio + Deposit_Ratio "
        "+ EntityEffects + TimeEffects"
    ),

    "Model 3 Nonlinear": (
        "Net_Fee_Margin ~ 1 + Log_Digital_Payment_Pressure "
        "+ Log_Digital_Payment_Pressure_Sq "
        "+ Capital_Adequacy_Proxy + Size + Credit_Risk "
        "+ Interest_Expense_Ratio + Deposit_Ratio "
        "+ EntityEffects + TimeEffects"
    )
}

results = {}

for name, formula in models.items():
    print("\n" + "=" * 90)
    print(name)
    print("=" * 90)

    model = PanelOLS.from_formula(formula, data=reg_df)
    result = model.fit(cov_type="clustered", cluster_entity=True)

    results[name] = result
    print(result)

# =====================================================
# 6. Helper functions for Word export
# =====================================================
def stars(p_value):
    if p_value < 0.01:
        return "***"
    elif p_value < 0.05:
        return "**"
    elif p_value < 0.10:
        return "*"
    else:
        return ""

def coef_and_se(result, variable):
    if variable in result.params.index:
        coef = result.params[variable]
        se = result.std_errors[variable]
        p = result.pvalues[variable]
        return f"{coef:.4f}{stars(p)}", f"({se:.4f})"
    return "", ""

def add_centered_paragraph(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(10)

# =====================================================
# 7. Create Word document
# =====================================================
doc = Document()

doc.add_heading("Advanced Regression Tests", level=1)

doc.add_paragraph(
    "This document reports additional fixed effects regression tests following the baseline "
    "results. The purpose is to investigate whether the insignificant baseline result may be "
    "explained by heterogeneous effects across bank size or by a non-linear relationship between "
    "digital payment pressure and banks' net fee margin."
)

doc.add_heading("Table X. Further Tests on Digital Payment Pressure", level=2)

var_map = {
    "Log Digital Payment Pressure": "Log_Digital_Payment_Pressure",
    "Digital Pressure × Large Bank": "Digital_x_LargeBank",
    "Log Digital Payment Pressure Squared": "Log_Digital_Payment_Pressure_Sq",
    "Capital Adequacy Proxy": "Capital_Adequacy_Proxy",
    "Size": "Size",
    "Credit Risk": "Credit_Risk",
    "Interest Expense Ratio": "Interest_Expense_Ratio",
    "Deposit Ratio": "Deposit_Ratio",
}

rows = [
    "Log Digital Payment Pressure",
    "",
    "Digital Pressure × Large Bank",
    "",
    "Log Digital Payment Pressure Squared",
    "",
    "Capital Adequacy Proxy",
    "",
    "Size",
    "",
    "Credit Risk",
    "",
    "Interest Expense Ratio",
    "",
    "Deposit Ratio",
    "",
    "Bank Fixed Effects",
    "Year Fixed Effects",
    "Clustered Standard Errors",
    "Observations",
    "Banks",
    "R-squared",
    "Within R-squared",
    "Robust F-stat p-value",
]

model_names = ["Model 1 Baseline", "Model 2 Interaction", "Model 3 Nonlinear"]

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"

headers = table.rows[0].cells
headers[0].text = "Variables"
headers[1].text = "Model 1 Baseline"
headers[2].text = "Model 2 Interaction"
headers[3].text = "Model 3 Nonlinear"

last_variable_label = None

for label in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = label

    if label in var_map:
        last_variable_label = label

    for i, model_name in enumerate(model_names, start=1):
        result = results[model_name]

        if label in var_map:
            coef, se = coef_and_se(result, var_map[label])
            add_centered_paragraph(row_cells[i], coef)

        elif label == "":
            if last_variable_label in var_map:
                coef, se = coef_and_se(result, var_map[last_variable_label])
                add_centered_paragraph(row_cells[i], se)
            else:
                add_centered_paragraph(row_cells[i], "")

        elif label == "Bank Fixed Effects":
            add_centered_paragraph(row_cells[i], "Yes")

        elif label == "Year Fixed Effects":
            add_centered_paragraph(row_cells[i], "Yes")

        elif label == "Clustered Standard Errors":
            add_centered_paragraph(row_cells[i], "Yes")

        elif label == "Observations":
            add_centered_paragraph(row_cells[i], str(int(result.nobs)))

        elif label == "Banks":
            n_banks = reg_df.index.get_level_values("gvkey").nunique()
            add_centered_paragraph(row_cells[i], str(n_banks))

        elif label == "R-squared":
            add_centered_paragraph(row_cells[i], f"{result.rsquared:.4f}")

        elif label == "Within R-squared":
            add_centered_paragraph(row_cells[i], f"{result.rsquared_within:.4f}")

        elif label == "Robust F-stat p-value":
            add_centered_paragraph(row_cells[i], f"{result.f_statistic_robust.pval:.4f}")

doc.add_paragraph(
    "Notes: Standard errors clustered at the bank level are reported in parentheses. "
    "*** p < 0.01, ** p < 0.05, * p < 0.10. "
    "The dependent variable is Net Fee Margin. "
    "All models include bank and year fixed effects."
)

doc.add_heading("Interpretation Guide", level=2)

doc.add_paragraph(
    "Model 1 reports the baseline specification using the log of digital payment pressure. "
    "Model 2 tests whether the effect of digital payment pressure differs between larger and "
    "smaller banks through an interaction term between digital pressure and the large bank dummy. "
    "Model 3 tests whether the relationship is non-linear by including the squared term of log "
    "digital payment pressure."
)

doc.add_paragraph(
    "For Model 2, a significant interaction term would suggest that the effect of digital payment "
    "pressure depends on bank size. For Model 3, a negative coefficient on the linear term and a "
    "positive coefficient on the squared term would be consistent with a U-shaped relationship, "
    "where digital payment pressure initially reduces banks' net fee margin but may become "
    "complementary at higher levels of digital development."
)

doc.save("advanced_regression_tests.docx")

print("\nWord file exported successfully: advanced_regression_tests.docx")