import os
import pandas as pd
import numpy as np
from linearmodels.panel import PanelOLS

# Optional Word export
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ModuleNotFoundError:
    DOCX_AVAILABLE = False
    print("python-docx is not installed. CSV outputs will still be created.")
    print("To export Word tables, run: python -m pip install python-docx")


# =====================================================
# 0. Settings
# =====================================================

INPUT_FILE = "bank_panel_with_digital_payment_pressure_and_Size_USD.csv"

OUTPUT_SAMPLE = "regression_sample_final_cleaned.csv"
OUTPUT_VALIDATION = "regression_validation_summary.csv"

MAIN_RESULTS_CSV = "main_regression_results_FINAL.csv"
ADVANCED_RESULTS_CSV = "advanced_regression_tests_FINAL.csv"

MAIN_DOCX = "main_regression_results_FINAL.docx"
ADVANCED_DOCX = "advanced_regression_tests_FINAL.docx"

APPLY_WINSORIZATION = True
WINSOR_LOWER = 0.01
WINSOR_UPPER = 0.99


# =====================================================
# 1. Load data
# =====================================================

if not os.path.exists(INPUT_FILE):
    raise FileNotFoundError(
        f"Cannot find {INPUT_FILE}. Make sure this file is in the same folder as this script."
    )

df = pd.read_csv(INPUT_FILE)

print("=" * 80)
print("Loaded data")
print("=" * 80)
print("Rows:", len(df))
print("Columns:", len(df.columns))


# =====================================================
# 2. Required columns check
# =====================================================

required_cols = [
    "gvkey",
    "fyear",
    "Net_Fee_Margin",
    "Digital_Payment_Pressure",
    "Capital_Adequacy_Proxy",
    "Size_USD",
    "pcl",
    "xint",
    "dptc",
    "at"
]

missing_cols = [c for c in required_cols if c not in df.columns]

if missing_cols:
    raise ValueError(f"Missing required columns: {missing_cols}")

print("\nAll required columns found.")


# =====================================================
# 3. Create regression variables
# =====================================================

# Convert relevant columns to numeric
numeric_cols = [
    "Net_Fee_Margin",
    "Digital_Payment_Pressure",
    "Capital_Adequacy_Proxy",
    "Size_USD",
    "pcl",
    "xint",
    "dptc",
    "at"
]

for col in numeric_cols:
    df[col] = pd.to_numeric(df[col], errors="coerce")

# Main bank-level controls
df["Credit_Risk"] = df["pcl"] / df["at"]
df["Interest_Expense_Ratio"] = df["xint"] / df["at"]
df["Deposit_Ratio"] = df["dptc"] / df["at"]

# Rescaled raw digital payment pressure:
# one unit = 10 percentage points of GDP
df["DPP_10pct_GDP"] = df["Digital_Payment_Pressure"] / 10.0

# Log-transformed digital payment pressure
df["Log_Digital_Payment_Pressure"] = np.log1p(df["Digital_Payment_Pressure"])

# Replace inf values with NaN
df = df.replace([np.inf, -np.inf], np.nan)


# =====================================================
# 4. Build final regression sample BEFORE Large_Bank median
# =====================================================

# Use a common sample for all main and advanced regressions
base_required_for_final_sample = [
    "gvkey",
    "fyear",
    "Net_Fee_Margin",
    "DPP_10pct_GDP",
    "Log_Digital_Payment_Pressure",
    "Capital_Adequacy_Proxy",
    "Size_USD",
    "Credit_Risk",
    "Interest_Expense_Ratio",
    "Deposit_Ratio"
]

sample_mask = df[base_required_for_final_sample].notna().all(axis=1)
reg_data = df.loc[sample_mask].copy()

print("\n" + "=" * 80)
print("Final regression sample before winsorization")
print("=" * 80)
print("Observations:", len(reg_data))
print("Banks:", reg_data["gvkey"].nunique())
print("Years:", sorted(reg_data["fyear"].unique().tolist()))

if len(reg_data) == 0:
    raise ValueError("Final regression sample has zero observations. Check missing values.")


# =====================================================
# 5. Winsorization within final regression sample
# =====================================================

def winsorize_series(s, lower=0.01, upper=0.99):
    lower_q = s.quantile(lower)
    upper_q = s.quantile(upper)
    return s.clip(lower=lower_q, upper=upper_q)

continuous_vars_to_winsorize = [
    "Net_Fee_Margin",
    "DPP_10pct_GDP",
    "Log_Digital_Payment_Pressure",
    "Capital_Adequacy_Proxy",
    "Size_USD",
    "Credit_Risk",
    "Interest_Expense_Ratio",
    "Deposit_Ratio"
]

if APPLY_WINSORIZATION:
    for col in continuous_vars_to_winsorize:
        reg_data[col] = winsorize_series(
            reg_data[col],
            lower=WINSOR_LOWER,
            upper=WINSOR_UPPER
        )

    print("\nWinsorization applied at 1st and 99th percentiles.")
else:
    print("\nWinsorization not applied.")


# =====================================================
# 6. Construct Large_Bank AFTER final sample and winsorization
# =====================================================

median_size_usd = reg_data["Size_USD"].median(skipna=True)

reg_data["Large_Bank"] = np.where(
    reg_data["Size_USD"] >= median_size_usd,
    1,
    0
)

# Interaction term
reg_data["Digital_x_LargeBank"] = (
    reg_data["Log_Digital_Payment_Pressure"] * reg_data["Large_Bank"]
)

print("\n" + "=" * 80)
print("Large_Bank construction")
print("=" * 80)
print("Median Size_USD:", median_size_usd)
print(reg_data["Large_Bank"].value_counts(dropna=False).sort_index())


# =====================================================
# 7. Mean-centre Log_DPP for nonlinear model
# =====================================================

mean_log_dpp = reg_data["Log_Digital_Payment_Pressure"].mean()

reg_data["Log_DPP_c"] = (
    reg_data["Log_Digital_Payment_Pressure"] - mean_log_dpp
)

reg_data["Log_DPP_c_Sq"] = reg_data["Log_DPP_c"] ** 2

print("\n" + "=" * 80)
print("Mean-centering for nonlinear model")
print("=" * 80)
print("Mean Log_Digital_Payment_Pressure:", mean_log_dpp)

corr_before = reg_data[["Log_Digital_Payment_Pressure"]].copy()
corr_before["Log_DPP_Sq_uncentered"] = reg_data["Log_Digital_Payment_Pressure"] ** 2
corr_uncentered = corr_before.corr().iloc[0, 1]

corr_after = reg_data[["Log_DPP_c", "Log_DPP_c_Sq"]].corr().iloc[0, 1]

print("Correlation: Log_DPP vs uncentered Log_DPP squared:", corr_uncentered)
print("Correlation: centered Log_DPP vs centered Log_DPP squared:", corr_after)


# =====================================================
# 8. Save final cleaned regression sample
# =====================================================

reg_data.to_csv(OUTPUT_SAMPLE, index=False)
print(f"\nSaved final regression sample: {OUTPUT_SAMPLE}")


# =====================================================
# 9. Validation summary
# =====================================================

validation_rows = []

validation_rows.append({
    "Item": "Input rows",
    "Value": len(df)
})

validation_rows.append({
    "Item": "Final regression observations",
    "Value": len(reg_data)
})

validation_rows.append({
    "Item": "Final regression banks",
    "Value": reg_data["gvkey"].nunique()
})

validation_rows.append({
    "Item": "Minimum year",
    "Value": reg_data["fyear"].min()
})

validation_rows.append({
    "Item": "Maximum year",
    "Value": reg_data["fyear"].max()
})

validation_rows.append({
    "Item": "Median Size_USD used for Large_Bank",
    "Value": median_size_usd
})

validation_rows.append({
    "Item": "Mean Log_DPP used for centering",
    "Value": mean_log_dpp
})

validation_rows.append({
    "Item": "Correlation before centering",
    "Value": corr_uncentered
})

validation_rows.append({
    "Item": "Correlation after centering",
    "Value": corr_after
})

validation_rows.append({
    "Item": "Winsorization applied",
    "Value": APPLY_WINSORIZATION
})

validation = pd.DataFrame(validation_rows)
validation.to_csv(OUTPUT_VALIDATION, index=False)

print(f"Saved validation summary: {OUTPUT_VALIDATION}")


# =====================================================
# 10. Set panel index
# =====================================================

reg_panel = reg_data.set_index(["gvkey", "fyear"])


# =====================================================
# 11. Model specifications
# =====================================================

main_models = {
    "Model 1": (
        "Net_Fee_Margin ~ 1 + DPP_10pct_GDP "
        "+ EntityEffects + TimeEffects"
    ),

    "Model 2": (
        "Net_Fee_Margin ~ 1 + DPP_10pct_GDP "
        "+ Capital_Adequacy_Proxy + Size_USD + Credit_Risk "
        "+ EntityEffects + TimeEffects"
    ),

    "Model 3": (
        "Net_Fee_Margin ~ 1 + DPP_10pct_GDP "
        "+ Capital_Adequacy_Proxy + Size_USD + Credit_Risk "
        "+ Interest_Expense_Ratio + Deposit_Ratio "
        "+ EntityEffects + TimeEffects"
    )
}

advanced_models = {
    "Model 4 Log Baseline": (
        "Net_Fee_Margin ~ 1 + Log_Digital_Payment_Pressure "
        "+ Capital_Adequacy_Proxy + Size_USD + Credit_Risk "
        "+ Interest_Expense_Ratio + Deposit_Ratio "
        "+ EntityEffects + TimeEffects"
    ),

    "Model 5 Interaction": (
        "Net_Fee_Margin ~ 1 + Log_Digital_Payment_Pressure "
        "+ Large_Bank "
        "+ Digital_x_LargeBank "
        "+ Capital_Adequacy_Proxy + Size_USD + Credit_Risk "
        "+ Interest_Expense_Ratio + Deposit_Ratio "
        "+ EntityEffects + TimeEffects"
    ),

    "Model 6 Nonlinear": (
        "Net_Fee_Margin ~ 1 + Log_DPP_c "
        "+ Log_DPP_c_Sq "
        "+ Capital_Adequacy_Proxy + Size_USD + Credit_Risk "
        "+ Interest_Expense_Ratio + Deposit_Ratio "
        "+ EntityEffects + TimeEffects"
    )
}


# =====================================================
# 12. Run models
# =====================================================

def run_models(model_dict, data):
    results = {}

    for name, formula in model_dict.items():
        print("\n" + "=" * 90)
        print(name)
        print("=" * 90)
        print(formula)

        model = PanelOLS.from_formula(
            formula,
            data=data,
            drop_absorbed=True,
            check_rank=False
        )

        result = model.fit(
            cov_type="clustered",
            cluster_entity=True
        )

        results[name] = result
        print(result)

    return results


main_results = run_models(main_models, reg_panel)
advanced_results = run_models(advanced_models, reg_panel)


# =====================================================
# 13. Export results to CSV
# =====================================================

def stars(p_value):
    if p_value < 0.01:
        return "***"
    elif p_value < 0.05:
        return "**"
    elif p_value < 0.10:
        return "*"
    else:
        return ""

def collect_results_to_dataframe(results, model_names, variable_order):
    rows = []

    for var_label, var_name in variable_order:
        row_coef = {"Variable": var_label}
        row_se = {"Variable": ""}

        for model_name in model_names:
            result = results[model_name]

            if var_name in result.params.index:
                coef = result.params[var_name]
                se = result.std_errors[var_name]
                p = result.pvalues[var_name]

                row_coef[model_name] = f"{coef:.6f}{stars(p)}"
                row_se[model_name] = f"({se:.6f})"
            else:
                row_coef[model_name] = ""
                row_se[model_name] = ""

        rows.append(row_coef)
        rows.append(row_se)

    summary_items = [
        ("Bank Fixed Effects", "Yes"),
        ("Year Fixed Effects", "Yes"),
        ("Clustered Standard Errors", "Yes"),
        ("Winsorization", "Yes" if APPLY_WINSORIZATION else "No"),
        ("Observations", None),
        ("Banks", None),
        ("R-squared", None),
        ("Within R-squared", None),
        ("Robust F-stat p-value", None),
    ]

    for item, fixed_value in summary_items:
        row = {"Variable": item}

        for model_name in model_names:
            result = results[model_name]

            if fixed_value is not None:
                row[model_name] = fixed_value
            elif item == "Observations":
                row[model_name] = str(int(result.nobs))
            elif item == "Banks":
                row[model_name] = str(reg_panel.index.get_level_values("gvkey").nunique())
            elif item == "R-squared":
                row[model_name] = f"{result.rsquared:.4f}"
            elif item == "Within R-squared":
                row[model_name] = f"{result.rsquared_within:.4f}"
            elif item == "Robust F-stat p-value":
                row[model_name] = f"{result.f_statistic_robust.pval:.4f}"

        rows.append(row)

    return pd.DataFrame(rows)


main_var_order = [
    ("Digital Payment Pressure (per 10 p.p. of GDP)", "DPP_10pct_GDP"),
    ("Capital Adequacy Proxy", "Capital_Adequacy_Proxy"),
    ("Size (USD)", "Size_USD"),
    ("Credit Risk", "Credit_Risk"),
    ("Interest Expense Ratio", "Interest_Expense_Ratio"),
    ("Deposit Ratio", "Deposit_Ratio"),
]

advanced_var_order = [
    ("Log Digital Payment Pressure", "Log_Digital_Payment_Pressure"),
    ("Large Bank", "Large_Bank"),
    ("Digital Pressure × Large Bank", "Digital_x_LargeBank"),
    ("Centered Log Digital Payment Pressure", "Log_DPP_c"),
    ("Centered Log Digital Payment Pressure Squared", "Log_DPP_c_Sq"),
    ("Capital Adequacy Proxy", "Capital_Adequacy_Proxy"),
    ("Size (USD)", "Size_USD"),
    ("Credit Risk", "Credit_Risk"),
    ("Interest Expense Ratio", "Interest_Expense_Ratio"),
    ("Deposit Ratio", "Deposit_Ratio"),
]

main_model_names = list(main_models.keys())
advanced_model_names = list(advanced_models.keys())

main_table_df = collect_results_to_dataframe(
    main_results,
    main_model_names,
    main_var_order
)

advanced_table_df = collect_results_to_dataframe(
    advanced_results,
    advanced_model_names,
    advanced_var_order
)

main_table_df.to_csv(MAIN_RESULTS_CSV, index=False)
advanced_table_df.to_csv(ADVANCED_RESULTS_CSV, index=False)

print(f"\nSaved main regression CSV: {MAIN_RESULTS_CSV}")
print(f"Saved advanced regression CSV: {ADVANCED_RESULTS_CSV}")


# =====================================================
# 14. Export Word documents
# =====================================================

def add_centered_paragraph(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.font.size = Pt(10)

def dataframe_to_word_table(doc, df_table):
    table = doc.add_table(rows=1, cols=len(df_table.columns))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, col in enumerate(df_table.columns):
        header_cells[i].text = col

    for _, row in df_table.iterrows():
        row_cells = table.add_row().cells
        for i, col in enumerate(df_table.columns):
            if i == 0:
                row_cells[i].text = str(row[col])
            else:
                add_centered_paragraph(row_cells[i], row[col])

    return table

if DOCX_AVAILABLE:

    # Main results Word
    main_doc = Document()
    main_doc.add_heading(
        "Main Regression Results: USD-adjusted Size and Rescaled Digital Pressure",
        level=1
    )

    main_doc.add_paragraph(
        "This table reports fixed effects regression results examining the relationship "
        "between digital payment pressure and banks' net fee margin. Digital payment pressure "
        "is measured as mobile and internet banking transaction value as a percentage of GDP, "
        "rescaled so that the coefficient represents a 10 percentage-point increase in "
        "GDP-equivalent pressure. Bank size is measured as the natural logarithm of total assets "
        "converted into USD. Continuous regression variables are winsorized at the 1st and 99th "
        "percentiles. Standard errors are clustered at the bank level."
    )

    main_doc.add_heading(
        "Table 1. Stepwise Fixed-Effects Models Using Rescaled Digital Payment Pressure",
        level=2
    )

    dataframe_to_word_table(main_doc, main_table_df)

    main_doc.add_paragraph(
        "Notes: Standard errors clustered at the bank level are reported in parentheses. "
        "*** p < 0.01, ** p < 0.05, * p < 0.10. "
        "The dependent variable is Net Fee Margin. "
        "All models include bank and year fixed effects. "
        "Size is measured as the natural logarithm of total assets converted into USD. "
        "Continuous regression variables are winsorized at the 1st and 99th percentiles."
    )

    main_doc.save(MAIN_DOCX)
    print(f"Saved main regression Word file: {MAIN_DOCX}")


    # Advanced results Word
    adv_doc = Document()
    adv_doc.add_heading(
        "Advanced Regression Tests: Log Transformation, Interaction and Mean-Centred Non-Linearity",
        level=1
    )

    adv_doc.add_paragraph(
        "This table reports refined fixed effects models using the final regression sample. "
        "The Large_Bank dummy is calculated only after excluding observations that cannot enter "
        "the regression, preventing missing digital pressure observations from influencing the "
        "median size threshold. Model 5 includes both the Large_Bank main effect and the interaction "
        "term. Model 6 uses a mean-centred log digital payment pressure variable before constructing "
        "the squared term to reduce multicollinearity between the linear and quadratic terms."
    )

    adv_doc.add_heading(
        "Table 2. Refined Fixed-Effects Models",
        level=2
    )

    dataframe_to_word_table(adv_doc, advanced_table_df)

    adv_doc.add_paragraph(
        "Notes: Standard errors clustered at the bank level are reported in parentheses. "
        "*** p < 0.01, ** p < 0.05, * p < 0.10. "
        "The dependent variable is Net Fee Margin. "
        "All models include bank and year fixed effects. "
        "Size is measured as the natural logarithm of total assets converted into USD. "
        "Continuous regression variables are winsorized at the 1st and 99th percentiles. "
        "In the nonlinear model, Log Digital Payment Pressure is mean-centred before squaring."
    )

    adv_doc.save(ADVANCED_DOCX)
    print(f"Saved advanced regression Word file: {ADVANCED_DOCX}")

else:
    print("\nWord export skipped because python-docx is not installed.")


# =====================================================
# 15. Final message
# =====================================================

print("\n" + "=" * 80)
print("DONE")
print("=" * 80)
print(f"Final cleaned sample: {OUTPUT_SAMPLE}")
print(f"Validation summary: {OUTPUT_VALIDATION}")
print(f"Main results CSV: {MAIN_RESULTS_CSV}")
print(f"Advanced results CSV: {ADVANCED_RESULTS_CSV}")

if DOCX_AVAILABLE:
    print(f"Main results Word: {MAIN_DOCX}")
    print(f"Advanced results Word: {ADVANCED_DOCX}")

print("\nImportant updates in this version:")
print("1. Large_Bank is calculated only after the final regression sample is formed.")
print("2. Model 5 includes both Large_Bank and Digital_x_LargeBank.")
print("3. Model 6 uses mean-centred Log_DPP and its squared term.")
print("4. Size is measured using Size_USD.")
print("5. Raw DPP is rescaled as DPP_10pct_GDP.")
print("6. Continuous variables are winsorized at the 1st and 99th percentiles.")