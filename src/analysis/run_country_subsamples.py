"""Estimate country-level robustness models using restricted thesis data.

The input sample is produced by src/analysis/run_main_regressions.py and is not
included in this public portfolio.
"""

from pathlib import Path
import pandas as pd
from linearmodels.panel import PanelOLS

# Optional Word export
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ModuleNotFoundError:
    DOCX_AVAILABLE = False
    print("python-docx is not installed. CSV outputs will still be created.")

# =====================================================
# 0. Settings
# =====================================================
ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "02_Data" / "04_Analysis_Output"
INPUT_FILE = OUTPUT_DIR / "ASEAN3_Regression_Sample_Cleaned.csv"
OUTPUT_CSV = OUTPUT_DIR / "ASEAN3_Country_Robustness.csv"
OUTPUT_DOCX = OUTPUT_DIR / "ASEAN3_Country_Robustness.docx"

# =====================================================
# 1. Load data
# =====================================================
if not INPUT_FILE.exists():
    raise FileNotFoundError(
        f"Cannot find {INPUT_FILE.name}. Generate it from the restricted main-analysis "
        "input; it is intentionally excluded from this public repository."
    )

reg_data = pd.read_csv(INPUT_FILE)

print("=" * 80)
print("Loaded cleaned regression data for Robustness Check")
print("=" * 80)

if "country" not in reg_data.columns:
    raise ValueError("The column 'country' is missing from the dataset. Make sure it is included in the upstream data merge.")

# =====================================================
# 2. Define Model Specification
# =====================================================
# Using the core Log Baseline specification (Model 4) for the robustness check
formula = (
    "Net_Fee_Margin ~ 1 + Log_Digital_Payment_Pressure "
    "+ Capital_Adequacy_Proxy + Size_USD + Credit_Risk "
    "+ Interest_Expense_Ratio + Deposit_Ratio "
    "+ EntityEffects"
)

target_countries = ["Indonesia", "Malaysia", "Thailand"]
results = {}

# =====================================================
# 3. Run Regressions by Country
# =====================================================
for country in target_countries:
    print(f"\nRunning Sub-sample Regression for: {country}")
    
    # Filter data for the specific country
    country_data = reg_data[reg_data["country"] == country].copy()
    
    if len(country_data) == 0:
        print(f"Warning: No data found for {country}. Skipping.")
        continue
        
    # Set panel index
    country_panel = country_data.set_index(["gvkey", "fyear"])
    
    # Fit model
    model = PanelOLS.from_formula(
        formula,
        data=country_panel,
        drop_absorbed=True,
        check_rank=False
    )
    
    result = model.fit(
        cov_type="clustered",
        cluster_entity=True
    )
    
    results[country] = result
    print(f"Observations: {result.nobs}, Banks: {country_panel.index.get_level_values('gvkey').nunique()}")

# =====================================================
# 4. Format Output Table
# =====================================================
def stars(p_value):
    if p_value < 0.01:
        return "***"
    elif p_value < 0.05:
        return "**"
    elif p_value < 0.10:
        return "*"
    else:
        return ""

var_order = [
    ("Log Digital Payment Pressure", "Log_Digital_Payment_Pressure"),
    ("Capital Adequacy Proxy", "Capital_Adequacy_Proxy"),
    ("Size (USD)", "Size_USD"),
    ("Credit Risk", "Credit_Risk"),
    ("Interest Expense Ratio", "Interest_Expense_Ratio"),
    ("Deposit Ratio", "Deposit_Ratio"),
]

rows = []
available_countries = list(results.keys())

for var_label, var_name in var_order:
    row_coef = {"Variable": var_label}
    row_se = {"Variable": ""}
    
    for country in available_countries:
        res = results[country]
        if var_name in res.params.index:
            coef = res.params[var_name]
            se = res.std_errors[var_name]
            p = res.pvalues[var_name]
            
            row_coef[country] = f"{coef:.6f}{stars(p)}"
            row_se[country] = f"({se:.6f})"
        else:
            row_coef[country] = ""
            row_se[country] = ""
            
    rows.append(row_coef)
    rows.append(row_se)

summary_items = [
    ("Bank Fixed Effects", "Yes"),
    ("Year Fixed Effects", "No"),  # 修正：由 Yes 改為 No
    ("Clustered Standard Errors", "Yes"),
    ("Observations", None),
    ("Banks", None),
    ("R-squared", None),
    ("Within R-squared", None),
]

for item, fixed_value in summary_items:
    row = {"Variable": item}
    for country in available_countries:
        res = results[country]
        if fixed_value is not None:
            row[country] = fixed_value
        elif item == "Observations":
            row[country] = str(int(res.nobs))
        elif item == "Banks":
            # Extract number of unique banks from the residual dataframe index
            row[country] = str(res.resids.index.get_level_values("gvkey").nunique())
        elif item == "R-squared":
            row[country] = f"{res.rsquared:.4f}"
        elif item == "Within R-squared":
            row[country] = f"{res.rsquared_within:.4f}"
            
    rows.append(row)

robustness_table_df = pd.DataFrame(rows)

# =====================================================
# 5. Export to CSV & Word
# =====================================================
robustness_table_df.to_csv(OUTPUT_CSV, index=False)
print(f"\nSaved Robustness regression CSV: {OUTPUT_CSV}")

if DOCX_AVAILABLE:
    def add_centered_paragraph(cell, text):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(text))
        run.font.size = Pt(10)

    doc = Document()
    doc.add_heading("Robustness Check: Country Sub-sample Analysis", level=1)
    
    doc.add_paragraph(
        "This table reports the fixed effects regression results for the ASEAN-3 pooled panel "
        "disaggregated into independent country sub-samples (Indonesia, Malaysia, and Thailand). "
        "The baseline log-transformed specification is utilized to isolate within-country effects. "
        "Note that year fixed effects are excluded in these sub-sample models because Log Digital "
        "Payment Pressure varies only over time within each country and would otherwise be absorbed "
        "by the time effects. This approach ensures the coefficient of the primary macro regressor "
        "can be statistically identified."
    )
    
    doc.add_heading("Table 3. Country Sub-sample Regressions", level=2)
    
    table = doc.add_table(rows=1, cols=len(robustness_table_df.columns))
    table.style = "Table Grid"
    
    header_cells = table.rows[0].cells
    for i, col in enumerate(robustness_table_df.columns):
        header_cells[i].text = col
        
    for _, row in robustness_table_df.iterrows():
        row_cells = table.add_row().cells
        for i, col in enumerate(robustness_table_df.columns):
            if i == 0:
                row_cells[i].text = str(row[col])
            else:
                add_centered_paragraph(row_cells[i], row[col])
                
    doc.add_paragraph(
        "Notes: Standard errors clustered at the bank level are reported in parentheses. "
        "*** p < 0.01, ** p < 0.05, * p < 0.10. The dependent variable is Net Fee Margin. "
        "All models include bank fixed effects but exclude year fixed effects due to perfect collinearity "
        "with the country-year digital payment variable. Size is measured as the natural logarithm "
        "of total assets converted into USD."
    )
    
    doc.save(OUTPUT_DOCX)
    print(f"Saved Robustness regression Word file: {OUTPUT_DOCX}")